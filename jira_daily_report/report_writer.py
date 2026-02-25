from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from jira_daily_report.date_utils import MonthWindow
from jira_daily_report.jira_client import DailyIssueSet


@dataclass(frozen=True)
class ReportPaths:
    spreadsheet_path: Path
    daily_summary_path: Path


def generate_reports(
    *,
    month_window: MonthWindow,
    daily_issue_set: DailyIssueSet,
    summaries_by_day: dict[date, str],
    jira_base_url: str,
    reports_dir: Path,
    spreadsheet_format: str,
) -> ReportPaths:
    reports_dir.mkdir(parents=True, exist_ok=True)

    working_days = _working_days_in_month(month_window)
    issue_hours_by_day = _issue_hours_by_day(daily_issue_set)
    month_label = month_window.start.strftime("%Y_%m")

    if spreadsheet_format == "csv":
        spreadsheet_path = reports_dir / f"{month_label}_jira_worklog.csv"
        _write_csv(
            path=spreadsheet_path,
            issue_hours_by_day=issue_hours_by_day,
            working_days=working_days,
            jira_base_url=jira_base_url,
        )
    else:
        spreadsheet_path = reports_dir / f"{month_label}_jira_worklog.xlsx"
        _write_xlsx(
            path=spreadsheet_path,
            issue_hours_by_day=issue_hours_by_day,
            working_days=working_days,
            jira_base_url=jira_base_url,
        )

    summary_path = reports_dir / f"{month_label}_daily_summary.txt"
    _write_daily_summary(path=summary_path, summaries_by_day=summaries_by_day)

    return ReportPaths(spreadsheet_path=spreadsheet_path, daily_summary_path=summary_path)


def _working_days_in_month(month_window: MonthWindow) -> list[date]:
    days: list[date] = []
    current = month_window.start
    while current <= month_window.end:
        if current.weekday() < 5:
            days.append(current)
        current += timedelta(days=1)
    return days


def _issue_hours_by_day(daily_issue_set: DailyIssueSet) -> dict[str, dict[date, float]]:
    result: dict[str, dict[date, float]] = {}
    for day, issue_seconds in daily_issue_set.by_day_issue_seconds.items():
        for issue_key, seconds in issue_seconds.items():
            by_day = result.setdefault(issue_key, {})
            by_day[day] = by_day.get(day, 0.0) + (seconds / 3600)
    return result


def _write_xlsx(
    *,
    path: Path,
    issue_hours_by_day: dict[str, dict[date, float]],
    working_days: list[date],
    jira_base_url: str,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Monthly Worklog"
    worksheet.freeze_panes = "B2"

    header = ["Jira Issue Link"] + [day.isoformat() for day in working_days]
    worksheet.append(header)

    for row_idx, issue_key in enumerate(sorted(issue_hours_by_day), start=2):
        issue_url = _issue_url(jira_base_url, issue_key)
        issue_cell = worksheet.cell(row=row_idx, column=1, value=issue_key)
        issue_cell.hyperlink = issue_url
        issue_cell.style = "Hyperlink"
        issue_cell.number_format = "@"

        for col_idx, day in enumerate(working_days, start=2):
            hours = issue_hours_by_day[issue_key].get(day)
            if hours is None:
                continue
            hour_cell = worksheet.cell(row=row_idx, column=col_idx, value=round(hours, 2))
            hour_cell.number_format = "0.00"

    worksheet.column_dimensions["A"].width = 25
    for col_idx in range(2, len(working_days) + 2):
        worksheet.column_dimensions[get_column_letter(col_idx)].width = 12

    workbook.save(path)


def _write_csv(
    *,
    path: Path,
    issue_hours_by_day: dict[str, dict[date, float]],
    working_days: list[date],
    jira_base_url: str,
) -> None:
    header = ["Jira Issue Link"] + [day.isoformat() for day in working_days]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)

        for issue_key in sorted(issue_hours_by_day):
            row = [_issue_url(jira_base_url, issue_key)]
            for day in working_days:
                hours = issue_hours_by_day[issue_key].get(day)
                row.append("" if hours is None else f"{round(hours, 2):g}")
            writer.writerow(row)


def _write_daily_summary(*, path: Path, summaries_by_day: dict[date, str]) -> None:
    lines = []
    for day in sorted(summaries_by_day):
        summary = summaries_by_day[day].strip()
        if not summary:
            continue
        lines.append(f"{day.isoformat()} - {summary}")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _issue_url(base_url: str, issue_key: str) -> str:
    return f"{base_url.rstrip('/')}/browse/{issue_key}"
